import os
import json
import uuid
import asyncio
import hashlib
import logging
import functools
import shutil

import grpc
import bm25s
import torch
import aiofiles
import aiofiles.os
import pandas as pd

from io import BytesIO
from datetime import datetime
from typing import List, Dict, Tuple, Any, Optional

from Stemmer import Stemmer
from pptx import Presentation
from docx import Document as DocxDocument
from docx.table import Table
from docx.oxml.ns import qn
from langchain_core.documents import Document as LangchainDocument
from sentence_transformers import SentenceTransformer
from qdrant_client import models as qmodels, AsyncQdrantClient


EMB_MODEL_NAME = "jinaai/jina-embeddings-v3"
COLLECTION = os.getenv("QDRANT_COLLECTION", "instructions")
DOC_DIR = os.getenv("DOC_DIR", "/app/data/instructions")
VECTOR_CACHE_DIR = os.getenv("VECTOR_CACHE_DIR", "/app/vector_db_cache")
HASH_FILE = os.path.join(VECTOR_CACHE_DIR, "hashes.json")
BM25_INDEX_DIR = os.path.join(VECTOR_CACHE_DIR, "bm25s_index")
BM25_METADATA_FILE = os.path.join(VECTOR_CACHE_DIR, "bm25s_metadata.json")


class RagManager:
    def __init__(self, logger: logging.Logger, qdrant_url: str = os.getenv("QDRANT_URL", "http://localhost:6334")):
        self.logger = logger
        self.qdrant_url = qdrant_url
        self.QDRANT_API_KEY = os.getenv("QDRANT_API_KEY", "")

        client_args = {
            "url": self.qdrant_url,
            "prefer_grpc": True,
            "timeout": 300
        }
        if self.QDRANT_API_KEY:
            self.logger.info("Using Qdrant Cloud")
            client_args["api_key"] = self.QDRANT_API_KEY
        else:
            client_args["url"] = "http://localhost:6334"
            self.logger.info("Using local Qdrant")
        self.client = AsyncQdrantClient(**client_args)

        emb_revision = os.getenv(
            "EMB_MODEL_REVISION",
            "ab036b023d30b4d1138c4c3bfa9f0c445ab455d6",
        )
        self.logger.info(
            f"Loading embedding model: {EMB_MODEL_NAME}@{emb_revision} (collection={COLLECTION})"
        )
        self.emb_model = SentenceTransformer(
            EMB_MODEL_NAME,
            device="cpu",
            trust_remote_code=True,
            revision=emb_revision,
            model_kwargs={
                "code_revision": emb_revision
            }
        )
        self.dim = self.emb_model.get_sentence_embedding_dimension()
        self.logger.info(f"Embedding dimension: {self.dim}")

        cpu_threads = max(1, int(os.getenv("TORCH_NUM_THREADS", os.cpu_count() or 1)))
        torch.set_num_threads(cpu_threads)
        try:
            torch.set_num_interop_threads(1)
        except RuntimeError:
            pass
        os.environ.setdefault("OMP_NUM_THREADS", str(cpu_threads))
        os.environ.setdefault("MKL_NUM_THREADS", str(cpu_threads))
        self._encode_lock: Optional[asyncio.Lock] = None
        self.logger.info(f"Encode will be serialized; torch threads={cpu_threads}")

        self.chunk_store: Dict[str, str] = {}
        self.list_idx_to_qdrant_id: Dict[int, str] = {}
        self.qdrant_id_to_list_idx: Dict[str, int] = {}
        self.bm25_retriever: Optional[bm25s.BM25] = None  
        try:
            self.stemmer = Stemmer("russian")
        except Exception as e:
            self.logger.error(f"Failed to initialize stemmer: {e}")
            self.stemmer = None

        os.makedirs(VECTOR_CACHE_DIR, exist_ok=True)
        os.makedirs(BM25_INDEX_DIR, exist_ok=True)

        self._bm25_meta_mtime_loaded = 0.0
        self._bm25_index_mtime_loaded = 0.0

    async def initialize(self, ensure_collection: bool = False):
        if self._encode_lock is None:
            self._encode_lock = asyncio.Lock()
        if ensure_collection:
            await self._ensure_collection()
        await self._load_bm25s_from_cache()
        if not self.bm25_retriever:
            self.logger.warning(
                "BM25s cache not found or outdated — run update_rag.py to rebuild from Qdrant"
            )
        elif not self.chunk_store or not self.list_idx_to_qdrant_id:
            self.logger.warning(
                "BM25s metadata not found or empty, will need rebuilding"
            )
            self.bm25_retriever = None

    async def _run_in_executor(self, func, *args):
        loop = asyncio.get_running_loop()
        return await loop.run_in_executor(None, func, *args)

    async def _ensure_collection(self):
        collection_exists = False
        current_payload_schema = {}
        desired_indexing_threshold = 50

        try:
            collection_info = await self.client.get_collection(COLLECTION)
            collection_exists = True
            self.logger.info(
                f"Collection '{COLLECTION}' already exists. Validating config"
            )
            vector_params = collection_info.config.params.vectors
            optimizer_params = collection_info.config.optimizer_config

            recreate_collection = False
            if vector_params.size != self.dim:
                self.logger.warning(
                    f"Vector size mismatch ({vector_params.size} vs {self.dim}). Recreating collection"
                )
                recreate_collection = True
            elif vector_params.distance != qmodels.Distance.COSINE:
                self.logger.warning(
                    f"Distance is not COSINE ('{vector_params.distance}'). Manual check recommended"
                )
            
            current_indexing_threshold = getattr(optimizer_params, 'indexing_threshold', None)
            if current_indexing_threshold is None or current_indexing_threshold != desired_indexing_threshold:
                self.logger.warning(
                    f"Indexing threshold mismatch or not set (current: {current_indexing_threshold}, desired: {desired_indexing_threshold}). Recreating collection."
                )
                recreate_collection = True

            if recreate_collection:
                await self.client.delete_collection(COLLECTION)
                collection_exists = False
            
            if hasattr(collection_info, 'payload_schema'):
                current_payload_schema = collection_info.payload_schema

        except grpc.RpcError as e:
            if e.code() == grpc.StatusCode.NOT_FOUND:
                self.logger.info(f"Collection '{COLLECTION}' not found")
                collection_exists = False
            else:
                self.logger.error(f"gRPC error checking collection: {e}")
                raise e
        except Exception as e:
            self.logger.error(f"Unexpected error checking collection: {e}")
            raise e

        if not collection_exists:
            try:
                self.logger.info(f"Creating collection '{COLLECTION}'")
                await self.client.create_collection(
                    collection_name=COLLECTION,
                    vectors_config=qmodels.VectorParams(
                        size=self.dim,
                        distance=qmodels.Distance.COSINE,
                    ),
                    optimizers_config=qmodels.OptimizersConfigDiff(
                        indexing_threshold=desired_indexing_threshold
                    )
                )
                collection_exists = True
                current_payload_schema = {}
            except Exception as create_exc:
                self.logger.error(
                    f"Failed to create collection '{COLLECTION}': {create_exc}"
                )
                raise create_exc

        if collection_exists:
            desired_indices = {
                "metadata.file": {
                    "field_name": "metadata.file",
                    "field_schema": qmodels.PayloadSchemaType.KEYWORD
                },
                "metadata.chunk_index": {
                    "field_name": "metadata.chunk_index",
                    "field_schema": qmodels.PayloadSchemaType.INTEGER
                },
                "page_content": {
                    "field_name": "page_content",
                    "field_schema": qmodels.PayloadSchemaType.TEXT
                }
            }

            for field, config in desired_indices.items():
                existing_schema_info = current_payload_schema.get(field)

                if not existing_schema_info:
                    self.logger.info(
                        f"Payload index for field '{field}' not found. Creating"
                    )
                    try:
                        await self.client.create_payload_index(
                            collection_name=COLLECTION,
                            field_name=config["field_name"],
                            field_schema=config["field_schema"]
                        )
                    except Exception as index_exc:
                        self.logger.error(
                            f"Failed to create payload index for '{field}': {index_exc}"
                        )
                elif hasattr(
                    existing_schema_info,
                    'data_type'
                ) and existing_schema_info.data_type != config["field_schema"]:
                    self.logger.warning(
                        f"Payload index type mismatch for '{field}': found {existing_schema_info.data_type}, expected {config['field_schema']}. Manual check needed"
                    )
                else:
                    self.logger.debug(
                        f"Payload index for field '{field}' already exists or type matches"
                    )

    async def _load_docx_async(self, file_path: str) -> List[str]:
        try:
            async with aiofiles.open(file_path, "rb") as f:
                content = await f.read()
            doc = await self._run_in_executor(DocxDocument, BytesIO(content))
            pages: List[str] = []
            current: List[str] = []

            for block in doc.element.body.iterchildren():
                if block.tag == qn("w:p"):
                    buffer = ""
                    for node in block.iter():
                        tag = node.tag
                        if tag == qn("w:t"):
                            buffer += node.text or ""
                        elif tag == qn("w:tab"):
                            buffer += "\t"
                        elif tag == qn("w:lastRenderedPageBreak") or (
                            tag == qn("w:br") and node.get(qn("w:type")) == "page"
                        ):
                            if buffer.strip():
                                current.append(buffer.strip())
                                buffer = ""
                            page_text = "\n".join(current).strip()
                            if page_text:
                                pages.append(page_text)
                            current = []
                    if buffer.strip():
                        current.append(buffer.strip())
                elif block.tag == qn("w:tbl"):
                    for row in Table(block, doc).rows:
                        cells, prev = [], None
                        for cell in row.cells:
                            text = cell.text.strip()
                            if text and text != prev:
                                cells.append(text)
                            prev = text
                        if cells:
                            current.append(" | ".join(cells))

            page_text = "\n".join(current).strip()
            if page_text:
                pages.append(page_text)

            return pages
        except Exception as e:
            self.logger.error(f"Error loading DOCX '{file_path}': {e}")
            return []

    async def _load_pptx_async(self, file_path: str) -> List[str]:
        try:
            async with aiofiles.open(file_path, "rb") as f:
                content = await f.read()
            pres = await self._run_in_executor(Presentation, BytesIO(content))
            slides_text = []
            for slide in pres.slides:
                texts_on_slide = []
                non_table_texts = []
                table_texts = []
                for shape in slide.shapes:
                    if shape.has_table:
                        table = shape.table
                        for row in table.rows:
                            for cell in row.cells:
                                if cell.text:
                                    table_texts.append(cell.text.strip())
                    elif hasattr(shape, "text"):
                        if shape.text:
                            top = int(getattr(shape, "top", 0))
                            left = int(getattr(shape, "left", 0))
                            non_table_texts.append((top, left, shape.text.strip()))

                if non_table_texts:
                    non_table_texts.sort(key=lambda x: (x[0], x[1]))
                    texts_on_slide.extend(t for _, _, t in non_table_texts)
                if table_texts:
                    texts_on_slide.extend(table_texts)

                slide_content = "\n".join(texts_on_slide)
                if slide_content:
                    slides_text.append(slide_content)
            return slides_text
        except Exception as e:
            self.logger.error(f"Error loading PPTX '{file_path}': {e}")
            return []

    async def _load_xlsx_async(self, file_path: str) -> List[str]:
        try:
            async with aiofiles.open(file_path, "rb") as f:
                content = await f.read()
            xls = await self._run_in_executor(pd.ExcelFile, BytesIO(content))
            sheets_text = []
            for sheet_name in xls.sheet_names:
                parse_func = functools.partial(
                    xls.parse,
                    sheet_name,
                    dtype=str
                )
                df = await self._run_in_executor(parse_func)
                df = df.fillna("")
                to_csv_func = functools.partial(
                    df.to_csv,
                    index=False,
                    header=True,
                    sep="\t"
                )
                sheet_content = await self._run_in_executor(to_csv_func)
                if sheet_content.strip():
                    sheets_text.append(sheet_content.strip())
            return sheets_text
        except Exception as e:
            self.logger.error(f"Error loading XLSX '{file_path}': {e}")
            return []

    async def _chunk_file(self, file_path: str) -> List[LangchainDocument]:
        ext = file_path.lower().split(".")[-1]
        raw_chunks = []
        if ext == "docx":
            raw_chunks = await self._load_docx_async(file_path)
        elif ext in ("pptx", "ppsx"):
            raw_chunks = await self._load_pptx_async(file_path)
        elif ext == "xlsx":
            raw_chunks = await self._load_xlsx_async(file_path)
        else:
            self.logger.warning(
                f"Unsupported file format: {file_path}"
            )
            return []

        docs = []
        try:
            stat_result = await aiofiles.os.stat(file_path)
            mtime = stat_result.st_mtime
            updated_iso = datetime.fromtimestamp(mtime).isoformat()
            base_name = os.path.basename(file_path)
            for i, chunk_text in enumerate(raw_chunks):
                if not chunk_text:
                    continue
                docs.append(LangchainDocument(
                    page_content=chunk_text,
                    metadata={
                        "file": base_name,
                        "chunk_index": i,
                        "source": file_path,
                        "updated_at": updated_iso
                    }
                ))
        except Exception as e:
            self.logger.error(
                f"Error creating documents for '{file_path}': {e}"
            )
        return docs

    async def _calculate_hash_async(self, file_path: str) -> Optional[str]:
        try:
            hasher = hashlib.md5()
            async with aiofiles.open(file_path, "rb") as f:
                while chunk := await f.read(8192):
                    hasher.update(chunk)
            return hasher.hexdigest()
        except FileNotFoundError:
            self.logger.warning(
                f"File not found while calculating hash: {file_path}"
            )
            return None
        except Exception as e:
            self.logger.error(
                f"Error calculating hash for '{file_path}': {e}"
            )
            return None

    async def sync(self):
        self.logger.info("Starting RAG index synchronization")
        try:
            hashes = {}
            try:
                if await aiofiles.os.path.exists(HASH_FILE):
                    async with aiofiles.open(HASH_FILE, "r") as f:
                        hashes = json.loads(await f.read())
            except json.JSONDecodeError:
                self.logger.warning(
                    f"Hash file {HASH_FILE} is corrupted"
                )
            current_files = set()
            files_to_reindex = {}
            tasks_chunking = []
            try:
                entries = await aiofiles.os.listdir(DOC_DIR)
            except FileNotFoundError:
                self.logger.error(f"Directory {DOC_DIR} not found")
                entries = []

            hash_tasks = []
            file_paths = []
            for filename in entries:
                full_path = os.path.join(DOC_DIR, filename)
                try:
                    if await aiofiles.os.path.isfile(full_path):
                        current_files.add(filename)
                        file_paths.append(full_path)
                        hash_tasks.append(
                            self._calculate_hash_async(full_path)
                        )
                except Exception as e:
                    self.logger.error(
                        f"Error accessing file {full_path}: {e}"
                    )
            calculated_hashes = await asyncio.gather(*hash_tasks)
            ids_to_delete = []
            for full_path, current_hash in zip(file_paths, calculated_hashes):
                filename = os.path.basename(full_path)
                if current_hash is None:
                    continue
                stored_hash = hashes.get(filename)
                if stored_hash != current_hash:
                    self.logger.info(f"File {filename} changed or new")
                    if stored_hash:
                        self.logger.info(
                            f"Scheduling deletion of old chunks for {filename}"
                        )
                        ids_to_delete.extend(
                            await self._get_ids_for_file(filename)
                        )
                        await self._delete_by_filename(filename)
                    files_to_reindex[full_path] = current_hash
                    tasks_chunking.append(
                        self._chunk_file(full_path)
                    )
                else:
                    hashes[filename] = current_hash

            removed_files = set(hashes.keys()) - current_files
            for filename in removed_files:
                self.logger.info(f"File {filename} removed; deleting points by filter")
                ids_to_delete.extend(await self._get_ids_for_file(filename))
                await self._delete_by_filename(filename)
                del hashes[filename]

            if ids_to_delete:
                self.logger.info(
                    f"Deleting {len(ids_to_delete)} old / deleted chunks from Qdrant"
                )
                await self.client.delete(
                    collection_name=COLLECTION,
                    points_selector=qmodels.PointIdsList(points=ids_to_delete),
                    wait=True
                )
            new_docs_nested = await asyncio.gather(*tasks_chunking)
            new_docs: List[LangchainDocument] = [
                doc for sublist in new_docs_nested for doc in sublist
            ]
            if new_docs:
                self.logger.info(
                    f"Received {len(new_docs)} new chunks for indexing"
                )
                texts = [d.page_content for d in new_docs]
                payloads = []
                for d in new_docs:
                    payload = {
                        "metadata": d.metadata.copy(),
                        "page_content": d.page_content,
                    }
                    payloads.append(payload)
                ids = [str(uuid.uuid4()) for _ in new_docs]

                self.logger.info("Calculating embeddings for new chunks")
                encode_call = functools.partial(
                    self.emb_model.encode,
                    normalize_embeddings=True,
                    batch_size=16,
                    task="retrieval.passage",
                    prompt_name="retrieval.passage",
                    convert_to_numpy=True
                )
                vectors = await self._run_in_executor(encode_call, texts)
                vectors_list = vectors.tolist()

                batch = qmodels.Batch(
                    ids=ids,
                    vectors=vectors_list,
                    payloads=payloads
                )
                response = await self.client.upsert(
                    collection_name=COLLECTION,
                    points=batch,
                    wait=True
                )
                self.logger.info(
                    f"Upsert in Qdrant completed: Status={response.status}"
                )
                for doc in new_docs:
                    file_path = doc.metadata.get("source")
                    if file_path in files_to_reindex:
                        hashes[
                            os.path.basename(file_path)
                        ] = files_to_reindex[file_path]
            else:
                self.logger.info("No new or changed files for indexing")

            async with aiofiles.open(HASH_FILE, "w") as f:
                await f.write(json.dumps(hashes, indent=4))
            await self._build_and_cache_bm25s()
            self.logger.info("RAG index synchronization completed")
        except Exception as e:
            self.logger.error(f"Error during RAG synchronization: {e}")

    async def _get_ids_for_file(self, filename: str) -> List[str]:
        try:
            ids = []
            for key in ("metadata.file", "file"):
                scroll_response, _ = await self.client.scroll(
                    collection_name=COLLECTION,
                    scroll_filter=qmodels.Filter(
                        must=[qmodels.FieldCondition(
                            key=key,
                            match=qmodels.MatchValue(value=filename)
                        )]
                    ),
                    limit=10000,
                    with_payload=False,
                    with_vectors=False
                )
                ids.extend([point.id for point in scroll_response])
            return ids
        except Exception as e:
            self.logger.error(f"Error getting IDs for file '{filename}': {e}")
            return []

    async def _delete_by_filename(self, filename: str):
        try:
            flt = qmodels.Filter(should=[
                qmodels.FieldCondition(key="metadata.file", match=qmodels.MatchValue(value=filename)),
                qmodels.FieldCondition(key="file", match=qmodels.MatchValue(value=filename)),
            ])
            await self.client.delete(
                collection_name=COLLECTION,
                points_selector=qmodels.FilterSelector(filter=flt),
                wait=True
            )
        except Exception as e:
            self.logger.error(f"Error deleting by filename '{filename}': {e}")

    async def _fetch_all_chunks_for_bm25s(self) -> Tuple[Dict[str, str], List[str], List[str]]:
        chunk_store: Dict[str, str] = {}
        ordered_qdrant_ids: List[str] = []
        ordered_texts: List[str] = []
        limit = 1000
        next_offset = None
        self.logger.info("Loading chunks from Qdrant for BM25s")
        retrieved_count = 0
        processed_points = 0

        while True:
            try:
                points, next_offset = await self.client.scroll(
                    collection_name=COLLECTION,
                    limit=limit,
                    offset=next_offset,
                    with_payload=qmodels.PayloadSelectorInclude(
                        include=["page_content"]
                    ),
                    with_vectors=False
                )
                if not points:
                    break
                
                processed_points += len(points)
                self.logger.debug(f"Fetched {len(points)} points from Qdrant. Total processed so far: {processed_points}.")

                for point in points:
                    text = None
                    point_id_str = str(point.id)
                    if point.payload:
                        self.logger.debug(f"Payload for point ID {point_id_str}: {point.payload}")
                        text = point.payload.get("page_content")

                    if text and isinstance(text, str):
                        if point_id_str not in chunk_store:
                            chunk_store[point_id_str] = text
                            ordered_qdrant_ids.append(point_id_str)
                            ordered_texts.append(text)
                            retrieved_count += 1
                        else:
                            self.logger.debug(f"Chunk ID {point_id_str} already in chunk_store. Skipping duplicate.")
                    else:
                        self.logger.warning(
                            f"Chunk with ID {point_id_str} doesn't contain valid text. Payload was: {point.payload}. Skipping."
                        )
                self.logger.debug(
                    f"After processing batch: chunk_store size: {len(chunk_store)}, retrieved_count: {retrieved_count}. Next offset: {next_offset}"
                )
                if next_offset is None:
                    break
            except Exception as e:
                self.logger.error(
                    f"Error loading chunks from Qdrant for BM25s: {e}"
                )
                break
        self.logger.info(
            f"Total {len(chunk_store)} unique chunks loaded for BM25s. {retrieved_count} new texts added to lists."
        )
        return chunk_store, ordered_qdrant_ids, ordered_texts

    async def _build_and_cache_bm25s(self):
        self.logger.info("Starting BM25s index build/update")
        chunk_store, ordered_qdrant_ids, ordered_texts = await self._fetch_all_chunks_for_bm25s()

        if not chunk_store or not ordered_texts:
            self.logger.warning("No data for BM25s index build")
            self.bm25_retriever = None
            self.chunk_store = {}
            self.list_idx_to_qdrant_id = {}
            if await aiofiles.os.path.exists(BM25_INDEX_DIR):
                await self._run_in_executor(shutil.rmtree, BM25_INDEX_DIR)
                await aiofiles.os.makedirs(BM25_INDEX_DIR, exist_ok=True)
            if await aiofiles.os.path.exists(BM25_METADATA_FILE):
                 await aiofiles.os.remove(BM25_METADATA_FILE)
            return

        self.chunk_store = chunk_store
        self.list_idx_to_qdrant_id = {
            i: q_id for i,
            q_id in enumerate(ordered_qdrant_ids)
        }
        self.qdrant_id_to_list_idx = {
            q_id: i for i,
            q_id in self.list_idx_to_qdrant_id.items()
        }
        try:
            self.logger.info(
                f"Indexing {len(ordered_texts)} documents with BM25s"
            )
            self.bm25_retriever = bm25s.BM25(backend="numba")

            if not self.stemmer:
                self.logger.error("Stemmer is not available. BM25s indexing will proceed without stemming.")
                tokenize_corpus_func = functools.partial(
                    bm25s.tokenize,
                    stopwords="ru"
                )
            else:
                tokenize_corpus_func = functools.partial(
                    bm25s.tokenize,
                    stemmer=self.stemmer,
                    stopwords="ru"
                )

            corpus_tokens = None
            try:
                corpus_tokens = await self._run_in_executor(tokenize_corpus_func, ordered_texts)
            except Exception as e_tokenize_executor:
                self.logger.error(f"Exception directly from _run_in_executor(tokenize_corpus_func): {e_tokenize_executor}")

            self.logger.info(f"Indexing {len(ordered_texts)} tokenized documents with BM25s.")
            await self._run_in_executor(
                self.bm25_retriever.index,
                corpus_tokens
            )
            self.logger.info(f"Saving BM25s index to {BM25_INDEX_DIR}")
            await self._run_in_executor(
                self.bm25_retriever.save,
                BM25_INDEX_DIR
            )
            metadata_cache = {
                "chunk_store": self.chunk_store,
                "list_idx_to_qdrant_id": self.list_idx_to_qdrant_id
            }
            dumps_with_indent = functools.partial(json.dumps, indent=4)
            metadata_json_string = await self._run_in_executor(
                dumps_with_indent,
                metadata_cache
            )
            async with aiofiles.open(BM25_METADATA_FILE, "w") as f:
                await f.write(metadata_json_string)
        except Exception as e:
            self.logger.error(f"Error building or caching BM25s: {e}")
            self.bm25_retriever = None

    async def _load_bm25s_from_cache(self):
        index_exists = await aiofiles.os.path.isdir(BM25_INDEX_DIR)
        metadata_exists = await aiofiles.os.path.exists(BM25_METADATA_FILE)

        if index_exists and metadata_exists:
            self.logger.info(f"Loading BM25s index from {BM25_INDEX_DIR}")
            try:
                load_with_mmap = functools.partial(
                    bm25s.BM25.load,
                    mmap=True
                )
                self.bm25_retriever = await self._run_in_executor(
                    load_with_mmap,
                    BM25_INDEX_DIR
                )
                self.bm25_retriever.backend = "numba"
                async with aiofiles.open(BM25_METADATA_FILE, "r") as f:
                    content = await f.read()
                metadata_cache = await self._run_in_executor(
                    json.loads,
                    content
                )
                self.chunk_store = metadata_cache.get("chunk_store", {})
                loaded_map = metadata_cache.get("list_idx_to_qdrant_id", {})
                self.list_idx_to_qdrant_id = {
                    int(k): v for k,
                    v in loaded_map.items()
                }
                self.qdrant_id_to_list_idx = {
                    v: k for k,
                    v in self.list_idx_to_qdrant_id.items()
                }
                try:
                    st_meta = await aiofiles.os.stat(BM25_METADATA_FILE)
                    st_index = await aiofiles.os.stat(BM25_INDEX_DIR)
                    self._bm25_meta_mtime_loaded = st_meta.st_mtime
                    self._bm25_index_mtime_loaded = st_index.st_mtime
                except Exception as e:
                    self.logger.warning(f"Error getting stats for BM25s files: {e}")
                if not self.chunk_store or not self.list_idx_to_qdrant_id:
                    self.logger.warning(
                        "Loaded BM25s metadata is empty. Index will be rebuilt"
                    )
                    self.bm25_retriever = None
                    self.chunk_store = {}
                    self.list_idx_to_qdrant_id = {}
                else:
                    self.logger.info(
                        f"Loaded BM25s metadata ({len(self.chunk_store)} chunks)"
                    )
            except FileNotFoundError:
                self.logger.info("BM25s cache not found")
                self.bm25_retriever = None
                self.chunk_store = {}
                self.list_idx_to_qdrant_id = {}
            except Exception as e:
                self.logger.error(f"Error loading BM25s cache: {e}")
                self.bm25_retriever = None
                self.chunk_store = {}
                self.list_idx_to_qdrant_id = {}
        else:
            self.logger.info("BM25s cache (index or metadata) not found")
            self.bm25_retriever = None
            self.chunk_store = {}
            self.list_idx_to_qdrant_id = {}

    async def _ensure_bm25_freshness(self):
        try:
            if not await aiofiles.os.path.exists(BM25_METADATA_FILE):
                return
            st_meta = await aiofiles.os.stat(BM25_METADATA_FILE)
            st_idx = await aiofiles.os.stat(BM25_INDEX_DIR) if await aiofiles.os.path.isdir(BM25_INDEX_DIR) else None
            newer = (st_meta.st_mtime > (self._bm25_meta_mtime_loaded or 0.0)) or (
                st_idx and st_idx.st_mtime > (self._bm25_index_mtime_loaded or 0.0)
            )
            if newer:
                self.logger.info("Detected updated BM25 cache on disk. Reloading...")
                await self._load_bm25s_from_cache()
        except Exception as e:
            self.logger.warning(f"BM25 freshness check failed: {e}")
            
    async def hybrid_search(
        self,
        query: str,
        k_vector=6,
        k_bm25=6,
        top_k=6
    ) -> List[Dict[str, Any]]:
        """Hybrid search without reranker (voice latency): vector∥BM25 → RRF → top_k."""
        await self._ensure_bm25_freshness()
        if self.bm25_retriever is None or not hasattr(
            self.bm25_retriever,
            'retrieve'
        ) or not self.list_idx_to_qdrant_id:
            self.logger.warning(
                "BM25s retriever isn't ready or no ID map. Performing only vector search"
            )
            return await self._vector_search_only(query, k_vector, top_k)
        vector_search_task = self._vector_search(query, k_vector)
        bm25_search_task = self._bm25s_search(query, k_bm25)
        vector_hits, bm25_results = await asyncio.gather(
            vector_search_task,
            bm25_search_task
        )
        self.logger.info(f"Vector search returned {len(vector_hits)} results")
        self.logger.info(f"BM25s search returned {len(bm25_results)} results")

        rrf_k = 60
        fusion_scores: Dict[str, float] = {}
        for rank, hit in enumerate(vector_hits):
            point_id = str(hit.id)
            fusion_scores[point_id] = fusion_scores.get(point_id, 0.0) + 1.0 / (rank + rrf_k)
        for rank, (point_id, _) in enumerate(bm25_results):
            if point_id in self.chunk_store:
                fusion_scores[point_id] = fusion_scores.get(point_id, 0.0) + 1.0 / (rank + rrf_k)

        fused_ranking = sorted(
            fusion_scores.items(),
            key=lambda item: item[1],
            reverse=True
        )[:top_k]
        self.logger.info(
            f"After RRF {len(fused_ranking)} candidates selected (no rerank)"
        )
        if not fused_ranking:
            self.logger.warning(
                "Hybrid search (RRF) didn't return any results"
            )
            return []

        candidate_ids = [pid for pid, _ in fused_ranking]
        candidates_points = await self._get_points_by_ids(candidate_ids)
        point_payloads = {
            str(p.id): p.payload for p in candidates_points if p.payload
        }
        final_results = []
        missing_texts = 0
        for point_id, rrf_score in fused_ranking:
            payload = point_payloads.get(point_id)
            text = payload.get("page_content") if payload else None
            if text:
                final_results.append({
                    "page_content": text,
                    "metadata": payload.get("metadata", {}),
                    "score": rrf_score,
                })
            else:
                missing_texts += 1
                self.logger.warning(
                    f"Text not found for RRF candidate with ID {point_id}"
                )
        if missing_texts > 0:
            self.logger.warning(
                f"Skipped {missing_texts} RRF candidates due to missing text"
            )
        self.logger.info(
            f"Hybrid search completed. Final number of results: {len(final_results)}"
        )
        return final_results

    async def _vector_search(
        self,
        query: str,
        k: int
    ) -> List[qmodels.ScoredPoint]:
        try:
            encode_query_call = functools.partial(
                self.emb_model.encode,
                normalize_embeddings=True,
                task="retrieval.query",
                prompt_name="retrieval.query",
                convert_to_numpy=True
            )
            if self._encode_lock is None:
                self._encode_lock = asyncio.Lock()
            async with self._encode_lock:
                q_emb_np = await self._run_in_executor(encode_query_call, query)
            if q_emb_np is None:
                self.logger.error("Vector search error: encode returned None")
                return []
            q_emb = q_emb_np.tolist()
            search_result = await self.client.search(
                collection_name=COLLECTION,
                query_vector=q_emb,
                limit=k,
                score_threshold=0.35,
                with_payload=False,
                with_vectors=False
            )
            return search_result
        except Exception as e:
            self.logger.error(f"Vector search error: {e}")
            return []

    async def _bm25s_search(
        self,
        query: str,
        k: int
    ) -> List[Tuple[str, float]]:
        if self.bm25_retriever is None or not self.list_idx_to_qdrant_id:
            self.logger.warning(
                "BM25s retriever isn't initialized or no ID map"
            )
            return []
        try:
            if not hasattr(bm25s, 'tokenize'):
                raise AttributeError(f"Global bm25s.tokenize method not found for version {bm25s.__version__}. Cannot tokenize query.")

            self.logger.debug(f"Tokenizing query for BM25s: '{query}'")
            if not self.stemmer:
                self.logger.warning("Stemmer not available. Tokenizing query without stemming.")
                tokenize_query_func = functools.partial(
                    bm25s.tokenize,
                    stopwords="ru"
                )
            else:
                tokenize_query_func = functools.partial(
                    bm25s.tokenize,
                    stemmer=self.stemmer,
                    stopwords="ru"
                )
            
            query_tokens = await self._run_in_executor(tokenize_query_func, query)

            retrieve_call = functools.partial(self.bm25_retriever.retrieve, k=k)
            results_indices, results_scores = await self._run_in_executor(
                retrieve_call,
                query_tokens
            )
            doc_indices = results_indices[0]
            scores = results_scores[0]
            bm25_ranking = []
            for list_idx, score in zip(doc_indices, scores):
                qdrant_id = self.list_idx_to_qdrant_id.get(int(list_idx))
                if qdrant_id and score > 0:
                    bm25_ranking.append((qdrant_id, float(score)))
            return bm25_ranking
        except Exception as e:
            self.logger.error(f"BM25s search error: {e}")
            return []

    async def _get_points_by_ids(
        self,
        ids: List[str]
    ) -> List[qmodels.PointStruct]:
        if not ids:
            return []
        try:
            points = await self.client.retrieve(
                collection_name=COLLECTION,
                ids=ids,
                with_payload=True,
                with_vectors=False
            )
            return points
        except Exception as e:
            self.logger.error(f"Error getting points by IDs: {e}")
            return []

    async def _vector_search_only(
        self,
        query: str,
        k: int,
        top_k: int
    ) -> List[Dict[str, Any]]:
        self.logger.debug("Executing only vector search (no BM25, no rerank)")
        vector_hits = await self._vector_search(query, k)
        if not vector_hits:
            return []
        candidate_ids = [str(hit.id) for hit in vector_hits[:top_k]]
        candidates_points = await self._get_points_by_ids(candidate_ids)
        point_payloads = {
            str(p.id): p.payload for p in candidates_points if p.payload
        }
        vector_scores = {str(h.id): h.score for h in vector_hits}
        results = []
        for point_id in candidate_ids:
            payload = point_payloads.get(point_id)
            text = payload.get("page_content") if payload else None
            if text:
                results.append({
                    "page_content": text,
                    "metadata": payload.get("metadata", {}),
                    "score": vector_scores.get(point_id, 0.0),
                })
        return results

    async def refresh_bm25_cache(self):
        await self._build_and_cache_bm25s()
        await self._load_bm25s_from_cache()